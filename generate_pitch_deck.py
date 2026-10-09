from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
import os

def create_pitch_deck():
    prs = Presentation()

    # Helpers
    def set_dark_background(slide):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(15, 23, 42) # Dark Slate

    def format_title(title_shape, text):
        title_shape.text = text
        for p in title_shape.text_frame.paragraphs:
            p.alignment = PP_ALIGN.LEFT
            for run in p.runs:
                run.font.color.rgb = RGBColor(52, 211, 153) # Emerald Green
                run.font.bold = True
                run.font.name = "Arial"

    def add_text_box(slide, left, top, width, height, text, color=RGBColor(255, 255, 255), size=20, bold=False):
        txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = text
        p.font.color.rgb = color
        p.font.size = Pt(size)
        p.font.bold = bold
        return tf

    # --------------------------------------------------
    # SLIDE 1: Problem Statement & Team (As Requested)
    # --------------------------------------------------
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    set_dark_background(slide)
    format_title(slide.shapes.title, "CyberGuard - AI Phishing Detection")

    # Team Intro
    add_text_box(slide, 0.5, 1.5, 9.0, 0.5, "TEAM:", RGBColor(52, 211, 153), 18, True)
    add_text_box(slide, 1.5, 1.5, 8.0, 0.5, "We are a team of cybersecurity enthusiasts bridging AI and web security.", RGBColor(226, 232, 240), 18, False)

    # Problem Statement
    add_text_box(slide, 0.5, 2.5, 9.0, 0.5, "THE PROBLEM:", RGBColor(248, 113, 113), 28, True)
    
    tf_prob = add_text_box(slide, 0.5, 3.2, 9.0, 3.0, "", RGBColor(226, 232, 240), 22, False)
    bullets = [
        "Phishing attacks account for 90% of all data breaches.",
        "Traditional blacklist scanners (like VirusTotal) fail against zero-day phishing URLs.",
        "Attackers use advanced evasion (typosquatting, obfuscation, URL shorteners).",
        "Everyday users have no real-time, explainable way to verify a suspicious link."
    ]
    for b in bullets:
        p = tf_prob.add_paragraph()
        p.text = f"• {b}"
        p.font.color.rgb = RGBColor(226, 232, 240)
        p.font.size = Pt(22)
        p.space_before = Pt(10)

    # --------------------------------------------------
    # SLIDE 2: The Solution
    # --------------------------------------------------
    slide2 = prs.slides.add_slide(prs.slide_layouts[5])
    set_dark_background(slide2)
    format_title(slide2.shapes.title, "The Solution: CyberGuard")

    add_text_box(slide2, 0.5, 1.5, 9.0, 0.5, "A real-time, Machine Learning-powered URL Analysis Engine.", RGBColor(255, 255, 255), 24, True)

    tf_sol = add_text_box(slide2, 0.5, 2.5, 9.0, 4.0, "", RGBColor(226, 232, 240), 22, False)
    bullets2 = [
        "Predictive AI: Trained on thousands of verified phishing and legitimate URLs.",
        "Contextual Intelligence: Extracts deep structural features (entropy, typosquatting).",
        "Instant Verdict: Processes and scores URLs in milliseconds via FastAPI.",
        "Beautiful UI: A sleek Next.js dashboard providing actionable explanations to users."
    ]
    for b in bullets2:
        p = tf_sol.add_paragraph()
        p.text = f"✓ {b}"
        p.font.color.rgb = RGBColor(52, 211, 153)
        p.font.size = Pt(22)
        p.space_before = Pt(14)

    # --------------------------------------------------
    # SLIDE 3: Architecture & Tech Stack
    # --------------------------------------------------
    slide3 = prs.slides.add_slide(prs.slide_layouts[5])
    set_dark_background(slide3)
    format_title(slide3.shapes.title, "How We Built It (Tech Stack)")

    tf_tech = add_text_box(slide3, 0.5, 2.0, 9.0, 4.0, "", RGBColor(226, 232, 240), 22, False)
    
    tech_stacks = [
        "Frontend: Next.js, React, TailwindCSS, Lucide Icons",
        "Backend API: FastAPI (Python), Uvicorn",
        "AI/ML Engine: Scikit-Learn (Logistic Regression / SVM)",
        "Data Pipeline: Pandas, TF-IDF Vectorization, Custom Feature Engineering"
    ]
    
    for stack in tech_stacks:
        p = tf_tech.add_paragraph()
        p.text = f"• {stack}"
        p.font.color.rgb = RGBColor(226, 232, 240)
        p.font.size = Pt(22)
        p.space_before = Pt(14)

    # --------------------------------------------------
    # SLIDE 4: Core Features / Differentiator
    # --------------------------------------------------
    slide4 = prs.slides.add_slide(prs.slide_layouts[5])
    set_dark_background(slide4)
    format_title(slide4.shapes.title, "Why CyberGuard Wins (Key Differentiators)")

    tf_diff = add_text_box(slide4, 0.5, 2.0, 9.0, 4.0, "", RGBColor(226, 232, 240), 22, False)
    diffs = [
        "Not Just a Blacklist: We catch brand new phishing links based on structural behavior.",
        "Privacy First: We don't track the user's browsing history.",
        "Explainable AI: Instead of just saying 'Dangerous', we tell the user WHY (e.g. 'Brand Impersonation Detected').",
        "Developer Friendly: Exposes a clean REST API for integration into email clients and browsers."
    ]
    for d in diffs:
        p = tf_diff.add_paragraph()
        p.text = f"🚀 {d}"
        p.font.color.rgb = RGBColor(226, 232, 240)
        p.font.size = Pt(22)
        p.space_before = Pt(14)

    # --------------------------------------------------
    # SLIDE 5: Business Impact & Roadmap
    # --------------------------------------------------
    slide5 = prs.slides.add_slide(prs.slide_layouts[5])
    set_dark_background(slide5)
    format_title(slide5.shapes.title, "Impact & Future Roadmap")

    add_text_box(slide5, 0.5, 1.5, 9.0, 0.5, "MARKET POTENTIAL:", RGBColor(96, 165, 250), 20, True)
    add_text_box(slide5, 0.5, 2.0, 9.0, 1.0, "B2B integration for enterprise email filters and B2C browser extensions.", RGBColor(226, 232, 240), 18, False)

    add_text_box(slide5, 0.5, 3.2, 9.0, 0.5, "ROADMAP (NEXT 6 MONTHS):", RGBColor(96, 165, 250), 20, True)
    tf_road = add_text_box(slide5, 0.5, 3.8, 9.0, 3.0, "", RGBColor(226, 232, 240), 18, False)
    
    roadmap = [
        "Phase 1: Deep learning models (Neural Networks) for higher accuracy.",
        "Phase 2: Chrome/Firefox Extension for real-time browser protection.",
        "Phase 3: Integration with global Threat Intelligence feeds (VirusTotal, URLhaus)."
    ]
    for r in roadmap:
        p = tf_road.add_paragraph()
        p.text = f"• {r}"
        p.font.color.rgb = RGBColor(226, 232, 240)
        p.font.size = Pt(20)
        p.space_before = Pt(10)

    # Save
    prs.save("CyberGuard_Hackathon_Pitch.pptx")

if __name__ == "__main__":
    create_pitch_deck()
    print("PPTX Created successfully.")
